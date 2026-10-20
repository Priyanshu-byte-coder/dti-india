"""Draw the human-validation sample and cut the coder form down to it.

    python scripts/make_validation_form.py [--n 200] [--n-keyword 70] [--seed 1] [--coders A B C]

Sample: stratified by NCO division x keyword flag (description names any digital term,
src/dti/baseline.py). Occupations that name digital work are ~6% of the corpus, so they are
oversampled to keep variation in the scores; the pilot occupations (used to develop the
prompt) are excluded. Inclusion probabilities and weights go to a sidecar CSV, not the form,
so coders see nothing but the occupation text.

Writes:
  data/interim/validation_<n>_seed<seed>.csv     sample with strata and weights
  forms/DTI_Validation_Form_<n>.xlsx             blank form (coder ID empty)
  forms/coders/DTI_Validation_Form_<n>_<ID>.xlsx one copy per coder ID, if --coders is given
"""
import argparse
import re
from copy import copy
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

from dti.baseline import keyword_baseline
from dti.config import CORPUS_CSV, FORM_XLSX, INTERIM, ROOT

FIRST, OLD_LAST = 3, 3452  # data rows in the full form


def allocate(sizes: pd.Series, n: int, minimum: int) -> pd.Series:
    """Proportional allocation with a floor, largest-remainder rounding, capped at stratum size."""
    quota = np.minimum(sizes, minimum).astype(int)
    while quota.sum() < n:
        room = sizes - quota
        if room.sum() == 0:
            break
        want = (n * sizes / sizes.sum()) - quota
        want[room == 0] = -np.inf
        quota[want.idxmax()] += 1
    return quota


def draw(n: int, n_keyword: int, seed: int) -> pd.DataFrame:
    corpus = pd.read_csv(CORPUS_CSV, dtype=str, keep_default_na=False)
    kw = keyword_baseline(corpus).set_index("nco2015")["kw_any"]
    corpus["kw_any"] = corpus["nco2015"].map(kw)
    pilot = set(pd.read_csv(INTERIM / "pilot_50_seed0.csv", dtype=str)["nco2015"])
    pool = corpus[~corpus["nco2015"].isin(pilot)]
    parts = []
    for flag, target in [(1, n_keyword), (0, n - n_keyword)]:
        sub = pool[pool["kw_any"] == flag]
        sizes = sub["division"].value_counts().sort_index()
        quota = allocate(sizes, target, minimum=2)
        for div, g in sub.groupby("division"):
            k = int(quota[div])
            s = g.sample(n=k, random_state=seed).copy()
            s["stratum"] = f"div{div}_kw{flag}"
            s["stratum_size"] = len(g)
            s["stratum_sample"] = k
            parts.append(s)
    out = pd.concat(parts).sort_values("nco2015").reset_index(drop=True)
    out["incl_prob"] = out["stratum_sample"] / out["stratum_size"]
    out["weight"] = 1 / out["incl_prob"]
    return out[["nco2015", "title", "division", "division_title", "kw_any", "stratum",
                "stratum_size", "stratum_sample", "incl_prob", "weight"]]


def cut_form(codes: list[str], out_path: Path, coder: str | None = None) -> None:
    wb = openpyxl.load_workbook(FORM_XLSX)
    ws = wb["Coding"]
    src_row = {str(ws.cell(r, 2).value): r for r in range(FIRST, OLD_LAST + 1)}
    keep = sorted(src_row[c] for c in codes)
    ncol = ws.max_column
    status = ws.cell(FIRST, 25).value  # Status formula of the first data row, used as the template
    # sources are ascending and never above their target row, so copying in order is safe
    for i, src in enumerate(keep):
        dst = FIRST + i
        for col in range(1, ncol + 1):
            a, b = ws.cell(src, col), ws.cell(dst, col)
            b.value = a.value
            if a.has_style:
                b._style = copy(a._style)
        ws.cell(dst, 1).value = i + 1
        ws.cell(dst, 25).value = re.sub(r"\b([A-Z]{1,2})3\b", rf"\g<1>{dst}", status)
        ws.row_dimensions[dst].height = ws.row_dimensions[src].height
    last = FIRST + len(keep) - 1
    ws.delete_rows(last + 1, OLD_LAST - last)

    def resize(ref: str) -> str:
        return " ".join(r.replace(str(OLD_LAST), str(last)) for r in str(ref).split())

    for dv in ws.data_validations.dataValidation:
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(resize(dv.sqref))
    # evidence must be words, not a number
    ev = DataValidation(type="custom", formula1="ISTEXT(K3)", allow_blank=True, showErrorMessage=True,
                        errorTitle="Evidence", error="Copy the exact words from the description, not a number.")
    ev.add(f"K{FIRST}:O{last}")
    ws.add_data_validation(ev)
    old_cf = ws.conditional_formatting
    ws.conditional_formatting = type(old_cf)()
    for rng, rules in old_cf._cf_rules.items():
        for rule in rules:
            ws.conditional_formatting.add(resize(rng.sqref), rule)
    ws.auto_filter.ref = f"A2:{openpyxl.utils.get_column_letter(ncol)}{last}"

    for row in wb["Progress"].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = c.value.replace(f"${OLD_LAST}", f"${last}")

    ins = wb["Instructions"]
    ins["A1"] = "Digital Task Intensity (DTI) — Human Validation Form, NCO-2015 occupations"
    ins["A2"] = (f"{len(keep)} occupations, a stratified sample of the NCO-2015 8-digit occupations. "
                 'Code each occupation on the "Coding" sheet. Yellow cells are for your answers.')
    ins["B17"] = ("Code blind and on your own: do not look at LLM scores or other coders' sheets, and do "
                  "not discuss scores with other coders until everyone has finished. Do not edit the grey columns.")
    ins["A19"], ins["B19"] = "9.", ("Judge each question separately. Q2–Q6 measure different things and "
                                    "often differ for the same job. Q7 evidence must be words copied from the "
                                    "description; a cell that holds only a number is rejected.")
    ins["A19"]._style, ins["B19"]._style = copy(ins["A18"]._style), copy(ins["B18"]._style)
    if coder:
        ins["C6"] = coder
    wb.save(out_path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=200)
    ap.add_argument("--n-keyword", type=int, default=70)
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--coders", nargs="*", default=[])
    a = ap.parse_args()

    sample = draw(a.n, a.n_keyword, a.seed)
    sample_csv = INTERIM / f"validation_{a.n}_seed{a.seed}.csv"
    sample.to_csv(sample_csv, index=False)
    blank = ROOT / "forms" / f"DTI_Validation_Form_{a.n}.xlsx"
    cut_form(list(sample["nco2015"]), blank)
    print(f"wrote {sample_csv} and {blank}")
    if a.coders:
        (ROOT / "forms" / "coders").mkdir(exist_ok=True)
        for c in a.coders:
            p = ROOT / "forms" / "coders" / f"DTI_Validation_Form_{a.n}_{c}.xlsx"
            cut_form(list(sample["nco2015"]), p, coder=c)
            print(f"wrote {p}")
    print(sample.groupby(["division", "kw_any"]).size().unstack(fill_value=0).assign(total=lambda d: d.sum(axis=1)))


if __name__ == "__main__":
    main()
